# to-pptx.py — assemble exported slide PNGs into a 16:9 PowerPoint (.pptx).
#   py scripts/to-pptx.py <deckDir> [out.pptx]
# Each build/export/*.png becomes one full-bleed 16:9 slide (pixel-perfect to the
# HTML design; text is baked into the image). The .pptx also imports cleanly into
# Google Slides (upload to Drive -> Open as Google Slides).
import sys, glob, os
from pptx import Presentation
from pptx.util import Inches

deck = os.path.abspath(sys.argv[1] if len(sys.argv) > 1 else ".")
pngs = sorted(glob.glob(os.path.join(deck, "build", "export", "*.png")))
if not pngs:
    sys.exit("no PNGs in build/export/ — run export.mjs first")
out = sys.argv[2] if len(sys.argv) > 2 else os.path.join(deck, "build", os.path.basename(deck) + ".pptx")

prs = Presentation()
prs.slide_width = Inches(13.333)   # 16:9 widescreen
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]
for p in pngs:
    s = prs.slides.add_slide(blank)
    s.shapes.add_picture(p, 0, 0, width=prs.slide_width, height=prs.slide_height)
prs.save(out)
print(f"wrote {out}  ({len(pngs)} slides, 13.333x7.5in 16:9)")
